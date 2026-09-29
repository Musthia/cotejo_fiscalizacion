"""Logica de carga de datos por boton y tabla.

Tablas:
- tableWidget_datos1 <- recursos/datos_1/...
- tableWidget_datos2 <- recursos/datos_2/...

Botones datos1:
- pushButton_datos_excel, pushButton_datos_csv, pushButton_datos_txt,
  pushButton_datos_db, pushButton_datos_acces, pushButton_nomb_archiv
Botones datos2: mismos nombres con sufijo _2.

Carpetas base (se crean si faltan):
- recursos/datos_1/{excel,csv,txt,db,access}
- recursos/datos_2/{excel,csv,txt,db,access}
- pushButton_nomb_archiv[_2] busca en recursos/datos_1 o datos_2 (cualquier tipo).
"""

import csv
import sqlite3
from pathlib import Path

from PySide6.QtWidgets import (
    QFileDialog,
    QInputDialog,
    QMessageBox,
    QTableWidget,
    QTableWidgetItem,
    QWidget,
)

BASE_DIR = Path(__file__).resolve().parent.parent
RECURSOS = BASE_DIR / "recursos"

TIPOS_VALIDOS = ("excel", "csv", "txt", "db", "access", "archivo")

# tipo logico -> subcarpeta dentro de datos_1 / datos_2
SUBCARPETA = {
    "excel": "excel",
    "csv": "csv",
    "txt": "txt",
    "db": "db",
    "access": "access",
    "archivo": "",  # raiz de datos_X (carga generica por extension)
}

FILTROS = {
    "excel": "Excel (*.xlsx *.xls)",
    "csv": "CSV (*.csv)",
    "txt": "Texto (*.txt)",
    "db": "SQLite (*.db *.sqlite *.sqlite3)",
    "access": "Access (*.mdb *.accdb)",
    "archivo": (
        "Todos los soportados (*.xlsx *.xls *.csv *.txt *.db *.sqlite *.sqlite3 *.mdb *.accdb);;"
        "Excel (*.xlsx *.xls);;CSV (*.csv);;Texto (*.txt);;"
        "SQLite (*.db *.sqlite *.sqlite3);;Access (*.mdb *.accdb);;Todos (*.*)"
    ),
}


def carpeta_para(tabla_num: int, tipo: str) -> Path:
    """Devuelve (y crea) la carpeta inicial del dialogo para tabla/tipo."""
    if tabla_num not in (1, 2):
        raise ValueError("tabla_num debe ser 1 o 2")
    if tipo not in TIPOS_VALIDOS:
        raise ValueError(f"tipo desconocido: {tipo}")
    carpeta = RECURSOS / f"datos_{tabla_num}" / SUBCARPETA[tipo]
    carpeta.mkdir(parents=True, exist_ok=True)
    return carpeta


def elegir_archivo(parent: QWidget | None, tabla_num: int, tipo: str) -> Path | None:
    """Abre QFileDialog en la carpeta correcta. Devuelve Path o None si cancela."""
    carpeta = carpeta_para(tabla_num, tipo)
    filtro = FILTROS[tipo]
    ruta, _ = QFileDialog.getOpenFileName(
        parent,
        f"Seleccionar {tipo.upper()} - Datos {tabla_num}",
        str(carpeta),
        filtro,
    )
    return Path(ruta) if ruta else None


def _leer_texto_delimitado(path: Path) -> tuple[list[str], list[list[str]]]:
    """Lee CSV/TXT detectando delimitador entre , ; TAB | . Primera fila = cabecera."""
    texto = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = path.read_text(encoding=encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if texto is None:
        raise ValueError(f"No se pudo decodificar {path.name}")

    lineas = [ln for ln in texto.splitlines() if ln.strip() != ""]
    if not lineas:
        return [], []

    muestra = "\n".join(lineas[:5])
    delimitador = ","
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=[",", ";", "\t", "|"])
        delimitador = dialecto.delimiter
    except csv.Error:
        # Heuristica simple: contar ocurrencias en la primera linea
        candidatos = {",": 0, ";": 0, "\t": 0, "|": 0}
        for sep in candidatos:
            candidatos[sep] = lineas[0].count(sep)
        mejor = max(candidatos, key=candidatos.get)
        if candidatos[mejor] > 0:
            delimitador = mejor
        elif "\t" in lineas[0]:
            delimitador = "\t"

    lector = csv.reader(lineas, delimiter=delimitador)
    filas = [[c.strip() for c in fila] for fila in lector if fila]
    cabecera, datos = filas[0], filas[1:]
    return cabecera, datos


def leer_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    return _leer_texto_delimitado(path)


def leer_txt(path: Path) -> tuple[list[str], list[list[str]]]:
    return _leer_texto_delimitado(path)


def leer_excel(path: Path) -> tuple[list[str], list[list[str]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "Para Excel instala: venv\\Scripts\\pip.exe install openpyxl"
        ) from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    wb.close()
    if not filas:
        return [], []
    cabecera = ["" if v is None else str(v) for v in filas[0]]
    datos = [["" if v is None else str(v) for v in fila] for fila in filas[1:]]
    return cabecera, datos


def _tablas_sqlite(path: Path) -> list[str]:
    con = sqlite3.connect(str(path))
    try:
        cur = con.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
        )
        return [r[0] for r in cur.fetchall()]
    finally:
        con.close()


def leer_sqlite(
    path: Path, parent: QWidget | None = None
) -> tuple[list[str], list[list[str]]]:
    tablas = _tablas_sqlite(path)
    if not tablas:
        raise ValueError(f"La base {path.name} no tiene tablas de usuario.")
    if len(tablas) == 1:
        tabla = tablas[0]
    else:
        tabla, ok = QInputDialog.getItem(
            parent,
            "Elegir tabla",
            f"{path.name} tiene {len(tablas)} tablas. Elige 1:",
            tablas,
            0,
            False,
        )
        if not ok or not tabla:
            raise InterruptedError("Seleccion de tabla cancelada.")
    con = sqlite3.connect(str(path))
    try:
        cur = con.execute(f'SELECT * FROM "{tabla}"')
        cabecera = [d[0] for d in cur.description] if cur.description else [tabla]
        datos = [["" if v is None else str(v) for v in fila] for fila in cur.fetchall()]
        return cabecera, datos
    finally:
        con.close()


def leer_access(
    path: Path, parent: QWidget | None = None
) -> tuple[list[str], list[list[str]]]:
    try:
        import pyodbc
    except ImportError as exc:
        raise RuntimeError(
            "Para Access instala: venv\\Scripts\\pip.exe install pyodbc "
            "(requiere driver 'Microsoft Access Driver (*.mdb, *.accdb)' en Windows)."
        ) from exc
    cadena = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={path};"
    try:
        con = pyodbc.connect(cadena)
    except Exception as exc:
        raise RuntimeError(
            f"No se pudo abrir {path.name} con el driver de Access. "
            "Verifica que el driver ODBC de 64 bits este instalado."
        ) from exc
    try:
        tablas = [t.table_name for t in con.cursor().tables(tableType="TABLE")]
        # Filtrar tablas de sistema de Access
        tablas = [t for t in tablas if not t.startswith("MSys")]
        if not tablas:
            raise ValueError(f"La base Access {path.name} no tiene tablas.")
        if len(tablas) == 1:
            tabla = tablas[0]
        else:
            tabla, ok = QInputDialog.getItem(
                parent,
                "Elegir tabla",
                f"{path.name} tiene {len(tablas)} tablas. Elige 1:",
                sorted(tablas),
                0,
                False,
            )
            if not ok or not tabla:
                raise InterruptedError("Seleccion de tabla cancelada.")
        cur = con.cursor()
        cur.execute(f'SELECT * FROM "{tabla}"')
        cabecera = [c[0] for c in cur.description]
        datos = [["" if v is None else str(v) for v in fila] for fila in cur.fetchall()]
        return cabecera, datos
    finally:
        con.close()


def leer_por_tipo(
    path: Path, tipo: str, parent: QWidget | None = None
) -> tuple[list[str], list[list[str]]]:
    """Despacha al lector correcto. tipo 'archivo' autodetecta por extension."""
    sufijo = path.suffix.lower()
    if tipo == "archivo":
        if sufijo in (".xlsx", ".xls"):
            return leer_excel(path)
        if sufijo == ".csv":
            return leer_csv(path)
        if sufijo == ".txt":
            return leer_txt(path)
        if sufijo in (".db", ".sqlite", ".sqlite3"):
            return leer_sqlite(path, parent)
        if sufijo in (".mdb", ".accdb"):
            return leer_access(path, parent)
        raise ValueError(f"Extension no soportada: {sufijo}")
    if tipo == "excel":
        return leer_excel(path)
    if tipo == "csv":
        return leer_csv(path)
    if tipo == "txt":
        return leer_txt(path)
    if tipo == "db":
        return leer_sqlite(path, parent)
    if tipo == "access":
        return leer_access(path, parent)
    raise ValueError(f"Tipo desconocido: {tipo}")


def mostrar_en_tabla(
    tabla: QTableWidget, cabecera: list[str], datos: list[list[str]]
) -> None:
    tabla.clear()
    tabla.setRowCount(0)
    tabla.setColumnCount(0)
    if not cabecera and not datos:
        return
    n_cols = len(cabecera) if cabecera else max(len(f) for f in datos)
    tabla.setColumnCount(n_cols)
    if cabecera:
        tabla.setHorizontalHeaderLabels(cabecera)
    tabla.setRowCount(len(datos))
    for i, fila in enumerate(datos):
        for j in range(n_cols):
            valor = fila[j] if j < len(fila) else ""
            tabla.setItem(i, j, QTableWidgetItem(valor))
    tabla.resizeColumnsToContents()


def cargar_en_tabla(
    parent: QWidget,
    tabla: QTableWidget,
    tabla_num: int,
    tipo: str,
) -> Path | None:
    """Flujo completo de un boton: dialogo en carpeta correcta -> leer -> mostrar.

    Devuelve la ruta cargada o None si se cancela / falla (muestra QMessageBox).
    """
    try:
        ruta = elegir_archivo(parent, tabla_num, tipo)
        if ruta is None:
            return None
        cabecera, datos = leer_por_tipo(ruta, tipo, parent)
        mostrar_en_tabla(tabla, cabecera, datos)
        tabla.setToolTip(f"{ruta.name} ({len(datos)} filas)")
        return ruta
    except InterruptedError:
        return None
    except Exception as exc:  # noqa: BLE001 - queremos mostrar el error al usuario
        QMessageBox.warning(parent, "Error al cargar", f"{type(exc).__name__}: {exc}")
        return None
